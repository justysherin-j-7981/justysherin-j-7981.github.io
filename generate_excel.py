"""
Generates token_data.xlsx with the following layout:

Columns A-D  : data rows
  A = Token
  B = Value_1
  C = Value_2
  D = Value_3

Columns E-G  : summary (distinct tokens, occurrence count, total row count)
  E = Value_1_count  – distinct Token values (unique list)
  F = Value_2_count  – how many times each distinct token appears in column A
  G = Value_3_count  – total number of data rows (same value on every summary row)

Run:
    python generate_excel.py
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import Counter

# ── sample data ──────────────────────────────────────────────────────────────
sample_data = [
    # (Token, Value_1, Value_2, Value_3)
    (456, 20, 10, 5),
    (456, 15,  8, 3),
    (789, 30, 12, 7),
    (123, 25, 11, 6),
    (789, 18,  9, 4),
    (123, 22, 13, 8),
    (456, 17, 10, 5),
    (321, 28, 14, 9),
]

# ── workbook / sheet setup ────────────────────────────────────────────────────
wb = Workbook()
ws = wb.active
ws.title = "Token Data"

# ── styles ────────────────────────────────────────────────────────────────────
header_font  = Font(bold=True, color="FFFFFF")
data_header_fill    = PatternFill("solid", fgColor="4472C4")   # blue
summary_header_fill = PatternFill("solid", fgColor="ED7D31")   # orange
center = Alignment(horizontal="center", vertical="center")
thin = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"),
)

def style_header(cell, fill):
    cell.font      = header_font
    cell.fill      = fill
    cell.alignment = center
    cell.border    = thin

def style_data(cell):
    cell.alignment = center
    cell.border    = thin

# ── write column headers (row 1) ──────────────────────────────────────────────
data_headers    = ["Token", "Value_1", "Value_2", "Value_3"]
summary_headers = ["Value_1_count", "Value_2_count", "Value_3_count"]

for col, name in enumerate(data_headers, start=1):
    cell = ws.cell(row=1, column=col, value=name)
    style_header(cell, data_header_fill)

for col, name in enumerate(summary_headers, start=5):
    cell = ws.cell(row=1, column=col, value=name)
    style_header(cell, summary_header_fill)

# ── write data rows (rows 2 … N+1) ───────────────────────────────────────────
for row_idx, (token, v1, v2, v3) in enumerate(sample_data, start=2):
    for col, value in enumerate([token, v1, v2, v3], start=1):
        cell = ws.cell(row=row_idx, column=col, value=value)
        style_data(cell)

# ── build summary (E, F, G) ───────────────────────────────────────────────────
tokens       = [row[0] for row in sample_data]
distinct     = sorted(set(tokens))          # unique tokens, sorted
token_counts = Counter(tokens)              # occurrences per token
total_rows   = len(sample_data)            # total data rows

for i, token in enumerate(distinct):
    row_idx = i + 2                        # start at row 2

    # E – distinct token value
    e_cell = ws.cell(row=row_idx, column=5, value=token)
    style_data(e_cell)

    # F – how many times this token appears in column A
    f_cell = ws.cell(row=row_idx, column=6, value=token_counts[token])
    style_data(f_cell)

    # G – total row count (same on every row)
    g_cell = ws.cell(row=row_idx, column=7, value=total_rows)
    style_data(g_cell)

# ── column widths ─────────────────────────────────────────────────────────────
col_widths = [12, 12, 12, 12, 18, 18, 18]
for i, width in enumerate(col_widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = width

# ── freeze header row ─────────────────────────────────────────────────────────
ws.freeze_panes = "A2"

# ── save ──────────────────────────────────────────────────────────────────────
output_file = "token_data.xlsx"
wb.save(output_file)
print(f"Saved: {output_file}")
